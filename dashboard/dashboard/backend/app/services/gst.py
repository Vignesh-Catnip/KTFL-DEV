"""GST master (GST_RATE.xlsx -> gst_tax_codes) and invoice GST % -> SAP tax code mapping."""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..models import GstTaxCode
from .textutil import snap_gst_rate


def parse_gst_description(desc: str):
    """
    'Input CGST+SGST @ 9%+9%'  -> (18.0,  9.0, CGST_SGST, INPUT)   tax code I3
    'Input IGST  18%'          -> (18.0, 18.0, IGST,      INPUT)   tax code I8
    'ND - Input IGST with 12%' -> (12.0, 12.0, IGST,      ND)      never picked
    """
    d = (desc or "").strip()
    u = d.upper()

    if u.startswith("TCS") or "TCS " in u:
        category = "TCS"
    elif u.startswith("ND"):
        category = "ND"
    elif u.startswith("RCM"):
        category = "RCM"
    elif u.startswith("ISD") or "ISD " in u:
        category = "ISD"
    elif "EXPORT" in u:
        category = "EXPORT"
    elif "IMPORT" in u:
        category = "IMPORT"
    elif u.startswith("OUTPUT"):
        category = "OUTPUT"
    elif "INPUT" in u:
        category = "INPUT"
    else:
        category = "OTHER"

    if "CGST" in u or "SGST" in u:
        nature = "CGST_SGST"
    elif "IGST" in u:
        nature = "IGST"
    else:
        nature = "OTHER"

    nums = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)\s*%", d)]
    if category == "TCS":
        rate = comp = (nums[-1] if nums else 0.0)          # TCS % is not added into the GST %
    elif nature == "CGST_SGST" and len(nums) >= 2:
        comp = round(nums[0], 2)
        rate = round(nums[0] + nums[1], 2)
    elif nature == "CGST_SGST" and len(nums) == 1:
        comp = round(nums[0], 2)
        rate = round(nums[0] * 2, 2) if nums[0] else 0.0
    else:
        rate = comp = round(nums[0], 2) if nums else 0.0
    return rate, comp, nature, category


def load_gst_master(db: Session, xlsx_path: str | Path, default_client: str | None = None) -> int:
    """Upsert every row of GST_RATE.xlsx. Returns the number of tax codes loaded."""
    p = Path(xlsx_path)
    if not p.exists():
        return 0
    client_default = default_client or get_settings().sap_client or "150"
    wb = load_workbook(p, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0
    hdr = [str(c or "").strip().lower() for c in rows[0]]
    idx = {name: hdr.index(name) for name in
           ("client", "procedure", "tax code", "tax type", "description") if name in hdr}
    if "tax code" not in idx:
        raise ValueError("GST master needs a 'Tax Code' column")

    n = 0
    for r in rows[1:]:
        if not r or not r[idx["tax code"]]:
            continue
        code = str(r[idx["tax code"]]).strip()
        desc = str(r[idx["description"]] or "").strip() if "description" in idx else ""
        client = str(r[idx["client"]] or client_default).strip() if "client" in idx else client_default
        proc = str(r[idx["procedure"]] or "").strip() if "procedure" in idx else ""
        ttype = str(r[idx["tax type"]] or "").strip() if "tax type" in idx else ""
        rate, comp, nature, category = parse_gst_description(desc)
        row = db.scalar(select(GstTaxCode).where(GstTaxCode.client == client, GstTaxCode.tax_code == code))
        if row is None:
            row = GstTaxCode(client=client, tax_code=code)
            db.add(row)
        row.procedure, row.tax_type, row.description = proc, ttype, desc
        row.rate, row.rate_component, row.nature, row.category = rate, comp, nature, category
        n += 1
    db.commit()
    return n


def pick_tax_code(db: Session, gst_total, gst_component, nature, is_import=False, is_rcm=False):
    """Invoice GST % -> purchase (input) tax code row. 9%+9% and 18% both find I3; IGST 18% finds I8.
    Only tax type 'V' / category INPUT (or IMPORT / RCM) is considered, so ND-, ISD, output and TCS
    codes can never be picked by accident."""
    if nature == "NIL" or gst_total is None:
        return None
    category = "IMPORT" if is_import else ("RCM" if is_rcm else "INPUT")
    rows = list(db.scalars(select(GstTaxCode).where(
        GstTaxCode.tax_type == "V", GstTaxCode.category == category, GstTaxCode.nature == nature,
    ).order_by(GstTaxCode.tax_code)))
    tol = 0.01
    cands = [r for r in rows if abs((r.rate or 0) - gst_total) <= tol]
    if not cands and gst_component is not None:
        cands = [r for r in rows if abs((r.rate_component or 0) - gst_component) <= tol]
    return cands[0] if cands else None


def compute_gst(h: dict) -> None:
    """Work out the GST % really charged on the invoice from the extracted amounts.
      gst_percent_component = per half (CGST 9 %)     gst_percent = total (9 + 9 = 18 %)"""
    sub, cgst, sgst, igst = h["subtotal"], h["cgst"], h["sgst"], h["igst"]
    tax_total = round(cgst + sgst + igst, 2) or h["total_tax"]

    if igst > 0 and (cgst + sgst) == 0:
        nature = "IGST"
    elif (cgst + sgst) > 0:
        nature = "CGST_SGST"
    elif tax_total == 0:
        nature = "NIL"
    else:
        nature = "IGST" if igst else "CGST_SGST"

    pct = comp = None
    if sub > 0 and tax_total > 0:
        pct = snap_gst_rate(round(tax_total / sub * 100, 2))
        if nature == "CGST_SGST":
            comp = (snap_gst_rate(round(max(cgst, sgst) / sub * 100, 2), component=True)
                    if (cgst or sgst) else round(pct / 2, 2))
        else:
            comp = pct
    elif h.get("gst_rate_printed", 0) > 0:
        r = snap_gst_rate(h["gst_rate_printed"])
        if nature == "CGST_SGST" and r and r <= 14:         # a printed "9%" on CGST+SGST = 9 per half
            comp, pct = r, snap_gst_rate(r * 2)
        else:
            pct = comp = r
    elif nature == "NIL":
        pct = comp = 0.0

    h["gst_percent"], h["gst_percent_component"], h["gst_nature"] = pct, comp, nature


def apply_tax_code(db: Session, h: dict, notes: list) -> None:
    """GST % on the invoice -> tax code from the master -> h['gst_tax_code'], then compare with SAP's."""
    compute_gst(h)
    pct, comp, nature = h["gst_percent"], h["gst_percent_component"], h["gst_nature"]
    h["gst_tax_code"] = ""
    h["gst_tax_code_desc"] = ""
    h["tax_code_match"] = None
    if pct is None:
        notes.append("GST % could not be computed - tax code not mapped")
        return
    shown = f"{comp}%+{comp}% (total {pct}%)" if nature == "CGST_SGST" else f"{pct}%"
    row = pick_tax_code(db, pct, comp, nature)
    if row:
        h["gst_tax_code"], h["gst_tax_code_desc"] = row.tax_code, row.description
        h["field_sources"]["gst_tax_code"] = "GST_MASTER"
        h["field_sources"]["gst_percent"] = "CALCULATED"
        if not h.get("tax_code"):
            h["tax_code"] = row.tax_code
        notes.append(f"GST {shown} {nature} -> tax code {row.tax_code} ({row.description})")
    else:
        notes.append(f"No input tax code in GST master for {shown} {nature}")

    if h.get("sap_tax_code") and h["gst_tax_code"]:
        h["tax_code_match"] = h["sap_tax_code"].upper() == h["gst_tax_code"].upper()
        if not h["tax_code_match"]:
            notes.append(f"Tax-code mismatch: invoice -> {h['gst_tax_code']} vs PO/GR -> {h['sap_tax_code']}")
