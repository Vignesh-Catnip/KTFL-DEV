"""
KTFL Invoice Extractor — LangChain + Ollama + PostgreSQL
=========================================================
Tables used:
  - invoices        : main invoice data
  - invoice_lines   : line items
  - invoice_master  : duplicate check (invoice_number UNIQUE)
  - audit_logs      : every action logged
  - processing_runs : RPA (unchanged)

Duplicate rule:
  invoice_master-ல invoice_number already இருந்தா → SKIP
  இல்லைன்னா → invoices push + invoice_master add
"""

import os
for k in ["HTTP_PROXY","HTTPS_PROXY","ALL_PROXY",
          "http_proxy","https_proxy","all_proxy"]:
    os.environ.pop(k, None)
os.environ["NO_PROXY"] = "localhost,127.0.0.1"
os.environ["no_proxy"] = "localhost,127.0.0.1"

import base64, json, re, sys, time
from datetime import datetime, timezone
from pathlib import Path

import fitz
import httpx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from langchain_core.messages import HumanMessage
from langchain_ollama import ChatOllama

from sqlalchemy import create_engine, text, select
from sqlalchemy.orm import (sessionmaker, DeclarativeBase,
                             Mapped, mapped_column, relationship)
from sqlalchemy import String, Text, Float, DateTime, ForeignKey, JSON, Integer

# ============================================================
# CONFIG
# ============================================================
PDF_FOLDER    = r"C:\Users\ckts00126\Desktop\KTFL-EXTRACTION\Input-pdf"
OLLAMA_URL    = "http://127.0.0.1:11434"
OLLAMA_MODEL  = "qwen3-vl:32b"
OUTPUT_FOLDER = r"C:\Users\ckts00126\Desktop\KTFL-EXTRACTION\Output-Excel"
DB_URL        = "postgresql+psycopg://appuser:appuser123@10.7.11.22:5432/miro_invoice"

MAX_PAGES    = 2
IMAGE_ZOOM   = 1.4
NUM_CTX      = 8192
NUM_PREDICT  = 2048
TIMEOUT      = 900
# ============================================================


# ── DB Models ────────────────────────────────────────────────

def now_utc(): return datetime.now(timezone.utc)

class Base(DeclarativeBase): pass

class Invoice(Base):
    __tablename__ = "invoices"
    id:             Mapped[int]      = mapped_column(Integer, primary_key=True)
    filename:       Mapped[str]      = mapped_column(String(255))
    vendor:         Mapped[str|None] = mapped_column(String(255))
    invoice_number: Mapped[str|None] = mapped_column(String(100), index=True)
    invoice_date:   Mapped[str|None] = mapped_column(String(30))
    po_number:      Mapped[str|None] = mapped_column(String(100))
    gstin:          Mapped[str|None] = mapped_column(String(50))
    hsn_sac:        Mapped[str|None] = mapped_column(String(50))
    vendor_code:    Mapped[str|None] = mapped_column(String(50))
    tax_code:       Mapped[str|None] = mapped_column(String(30))
    business_place: Mapped[str|None] = mapped_column(String(100))
    currency:       Mapped[str]      = mapped_column(String(10), default="INR")
    subtotal:       Mapped[float]    = mapped_column(Float, default=0)
    cgst:           Mapped[float]    = mapped_column(Float, default=0)
    sgst:           Mapped[float]    = mapped_column(Float, default=0)
    igst:           Mapped[float]    = mapped_column(Float, default=0)
    tax:            Mapped[float]    = mapped_column(Float, default=0)
    total:          Mapped[float]    = mapped_column(Float, default=0)
    confidence:     Mapped[float]    = mapped_column(Float, default=0)
    grn_number:     Mapped[str|None] = mapped_column(String(100))
    handwritten_references: Mapped[str|None] = mapped_column(Text)
    status:         Mapped[str]      = mapped_column(String(40), default="REVIEW", index=True)
    error_message:  Mapped[str|None] = mapped_column(Text)
    created_at:     Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    updated_at:     Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc, onupdate=now_utc)
    lines = relationship("InvoiceLine", back_populates="invoice", cascade="all, delete-orphan")

class InvoiceLine(Base):
    __tablename__ = "invoice_lines"
    id:         Mapped[int]   = mapped_column(primary_key=True)
    invoice_id: Mapped[int]   = mapped_column(ForeignKey("invoices.id", ondelete="CASCADE"))
    item:       Mapped[str]   = mapped_column(String(255))
    quantity:   Mapped[float] = mapped_column(Float, default=0)
    rate:       Mapped[float] = mapped_column(Float, default=0)
    tax_code:   Mapped[str]   = mapped_column(String(30), default="")
    amount:     Mapped[float] = mapped_column(Float, default=0)
    invoice = relationship("Invoice", back_populates="lines")

class InvoiceMaster(Base):
    __tablename__ = "invoice_master"
    id:             Mapped[int]      = mapped_column(primary_key=True)
    invoice_number: Mapped[str]      = mapped_column(String(100), unique=True)
    vendor:         Mapped[str|None] = mapped_column(String(255))
    vendor_code:    Mapped[str|None] = mapped_column(String(50))
    po_number:      Mapped[str|None] = mapped_column(String(100))
    first_seen_date: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)
    invoice_id:     Mapped[int|None] = mapped_column(ForeignKey("invoices.id", ondelete="SET NULL"))

class AuditLog(Base):
    __tablename__ = "audit_logs"
    id:         Mapped[int]      = mapped_column(primary_key=True)
    invoice_id: Mapped[int|None] = mapped_column(ForeignKey("invoices.id", ondelete="SET NULL"))
    action:     Mapped[str]      = mapped_column(String(100))
    detail:     Mapped[dict]     = mapped_column(JSON, default=dict)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=now_utc)


# ── DB Setup ─────────────────────────────────────────────────

def setup_db():
    engine  = create_engine(DB_URL)
    Session = sessionmaker(bind=engine)

    # Create tables if not exist
    Base.metadata.create_all(engine)

    # Auto-add new columns to existing tables
    migrations = [
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS gstin VARCHAR(50)",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS hsn_sac VARCHAR(50)",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS vendor_code VARCHAR(50)",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS cgst FLOAT DEFAULT 0",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS sgst FLOAT DEFAULT 0",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS igst FLOAT DEFAULT 0",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS grn_number VARCHAR(100)",
        "ALTER TABLE invoices ADD COLUMN IF NOT EXISTS handwritten_references TEXT",
        "ALTER TABLE invoice_lines ADD COLUMN IF NOT EXISTS rate FLOAT DEFAULT 0",
    ]
    with engine.connect() as conn:
        for sql in migrations:
            try: conn.execute(text(sql))
            except: pass
        conn.commit()

    print("  DB ready ✅")
    return engine, Session


# ── Duplicate Check ───────────────────────────────────────────

def is_duplicate(session, invoice_number: str) -> tuple[bool, str]:
    """
    invoice_master-ல invoice_number இருந்தா → DUPLICATE
    இல்லைன்னா → NEW
    """
    if not invoice_number:
        return False, "No invoice number — skip duplicate check"

    existing = session.scalar(
        select(InvoiceMaster).where(
            InvoiceMaster.invoice_number == invoice_number
        )
    )

    if existing:
        return True, (
            f"DUPLICATE: invoice_number '{invoice_number}' "
            f"already in invoice_master (first seen: {existing.first_seen_date.date()}, "
            f"vendor: {existing.vendor or '?'})"
        )
    return False, "New invoice — OK to push"


# ── DB Push ───────────────────────────────────────────────────

def push_to_db(session, header: dict, lines: list, pdf_name: str) -> int:
    """
    1. invoices table-ல insert
    2. invoice_lines table-ல insert
    3. invoice_master table-ல add (duplicate prevention)
    4. audit_logs table-ல log
    """

    # 1. Invoice
    inv = Invoice(
        filename       = pdf_name,
        vendor         = header["vendor"],
        invoice_number = header["invoice_number"],
        invoice_date   = header["invoice_date"],
        po_number      = header["po_number"],
        gstin          = header["gstin"],
        hsn_sac        = header["hsn_sac"],
        vendor_code    = header["vendor_code"],
        tax_code       = "",
        business_place = "",
        currency       = header["currency"],
        subtotal       = header["subtotal"],
        cgst           = header["cgst"],
        sgst           = header["sgst"],
        igst           = header["igst"],
        tax            = header["total_tax"],
        total          = header["total"],
        confidence     = header["confidence"] / 100,
        grn_number     = header["grn_number"],
        handwritten_references = header["handwritten_references"],
        status         = "REVIEW",
    )
    session.add(inv)
    session.flush()  # get inv.id

    # 2. Line items
    for l in lines:
        session.add(InvoiceLine(
            invoice_id = inv.id,
            item       = l["item"],
            quantity   = l["quantity"],
            rate       = l["rate"],
            tax_code   = l["hsn_sac"],
            amount     = l["amount"],
        ))

    # 3. Invoice master (duplicate prevention)
    session.add(InvoiceMaster(
        invoice_number  = header["invoice_number"],
        vendor          = header["vendor"],
        vendor_code     = header["vendor_code"],
        po_number       = header["po_number"],
        invoice_id      = inv.id,
    ))

    # 4. Audit log
    session.add(AuditLog(
        invoice_id = inv.id,
        action     = "EXTRACTED_AND_INSERTED",
        detail     = {
            "source"    : "invoice_extractor_db.py",
            "model"     : OLLAMA_MODEL,
            "file"      : pdf_name,
            "confidence": header["confidence"],
            "total"     : header["total"],
            "grn"       : header["grn_number"],
        }
    ))

    session.commit()
    return inv.id


# ── Extraction Prompt ─────────────────────────────────────────

PROMPT = r"""
You are an Indian invoice extraction engine.
Read the invoice IMAGE. Return ONLY valid JSON. No markdown. No explanation.

AMOUNT RULES:
  ₹94,540.14  → 94540.14
  ₹6,19,763.14 → 619763.14
  Remove ₹ and commas. Keep decimal point exactly.

GRN NUMBER:
  Look for receiving stamp: "GRN", "Inward No", "INWARD NO", "MRN"
  Extract ONLY that stamp number. NOT vehicle/challan/e-way bill numbers.

LINE ITEMS:
  Extract ONLY invoice line items with amounts. Remove duplicates.
  Apply invoice-level HSN/SAC to all items if only one HSN shown.

Return this JSON:
{
  "vendor": "",
  "invoice_number": "",
  "invoice_date": "",
  "po_number": "",
  "gstin": "",
  "hsn_sac": "",
  "vendor_code": "",
  "currency": "INR",
  "subtotal": 0.0,
  "cgst": 0.0,
  "sgst": 0.0,
  "igst": 0.0,
  "total_tax": 0.0,
  "total": 0.0,
  "grn_number": "",
  "handwritten_references": "",
  "confidence": 0.0,
  "lines": [
    {"item": "", "quantity": 0.0, "rate": 0.0, "hsn_sac": "", "amount": 0.0}
  ]
}
"""


# ── Helpers ───────────────────────────────────────────────────

def to_float(v) -> float:
    if v is None: return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = re.sub(r"[^0-9.\-]", "", str(v).replace(",","").replace("₹",""))
    try: return float(s) if s else 0.0
    except: return 0.0

def clean(v) -> str:
    return str(v).strip() if v else ""

def extract_json(raw) -> dict:
    if isinstance(raw, list):
        raw = "".join(str(x.get("text","") if isinstance(x,dict) else x) for x in raw)
    text = str(raw).strip()
    text = re.sub(r"<think>.*?</think>","",text,flags=re.DOTALL|re.IGNORECASE).strip()
    text = re.sub(r"^```(?:json)?","",text,flags=re.IGNORECASE)
    text = re.sub(r"```$","",text).strip()
    try: return json.loads(text)
    except: pass
    s, e = text.find("{"), text.rfind("}")
    if s>=0 and e>s:
        try: return json.loads(text[s:e+1])
        except: pass
    raise ValueError("No valid JSON in response")

def pdf_to_images(path: Path):
    doc  = fitz.open(str(path))
    imgs = []
    zoom = fitz.Matrix(IMAGE_ZOOM, IMAGE_ZOOM)
    for i in range(min(doc.page_count, MAX_PAGES)):
        pix = doc[i].get_pixmap(matrix=zoom, alpha=False)
        imgs.append(base64.b64encode(pix.tobytes("png")).decode())
    doc.close()
    return imgs

def get_pdf_text(path: Path) -> str:
    doc   = fitz.open(str(path))
    parts = []
    for i, page in enumerate(doc, 1):
        t = page.get_text("text")
        if t: parts.append(f"--- PAGE {i} ---\n{t[:2000]}")
    doc.close()
    return "\n".join(parts)

def fix_amounts(subtotal, cgst, sgst, igst, total_tax, total):
    if total > 0 and total_tax > total:
        candidate = round(total_tax / 100, 2)
        if abs(subtotal + candidate - total) < total * 0.02:
            total_tax = candidate
            cgst = round(cgst/100, 2) if cgst else 0
            sgst = round(sgst/100, 2) if sgst else 0
            igst = round(igst/100, 2) if igst else 0
    if subtotal > 0 and total_tax >= 0:
        expected = subtotal + total_tax
        if expected > 0 and total > expected * 10:
            total = round(total / 100, 2)
    if total_tax == 0:
        total_tax = round(cgst + sgst + igst, 2)
    if subtotal == 0 and total > 0:
        subtotal = round(total - total_tax, 2)
    if total == 0:
        total = round(subtotal + total_tax, 2)
    return subtotal, cgst, sgst, igst, total_tax, total

def create_llm():
    return ChatOllama(
        model=OLLAMA_MODEL,
        base_url=OLLAMA_URL,
        temperature=0,
        num_ctx=NUM_CTX,
        num_predict=NUM_PREDICT,
        format="json",
        client_kwargs={"trust_env": False},
    )

def invoke(llm, images, pdf_text="") -> dict:
    try:
        content = [{"type":"text","text":PROMPT}]
        if pdf_text:
            content.append({"type":"text","text":"\nOCR:\n"+pdf_text[:2000]})
        for img in images:
            content.append({"type":"image_url",
                "image_url":f"data:image/png;base64,{img}"})
        resp = llm.invoke([HumanMessage(content=content)])
        return extract_json(resp.content)
    except Exception as e:
        print(f"     LangChain failed: {e} — raw Ollama...")

    body = PROMPT
    if pdf_text: body += "\nOCR:\n" + pdf_text[:2000]
    payload = {
        "model": OLLAMA_MODEL,
        "messages": [{"role":"user","content":body,"images":images}],
        "stream": False, "format": "json", "think": False,
        "options": {"temperature":0,"num_ctx":NUM_CTX,"num_predict":NUM_PREDICT},
    }
    r = httpx.post(f"{OLLAMA_URL}/api/chat",
                   json=payload, timeout=TIMEOUT, trust_env=False)
    r.raise_for_status()
    msg = r.json().get("message",{})
    raw = (msg.get("content") or msg.get("thinking") or "").strip()
    if not raw: raise ValueError("Ollama returned empty response")
    return extract_json(raw)

def normalize(data: dict, filename: str):
    if not isinstance(data, dict): data = {}
    subtotal  = to_float(data.get("subtotal"))
    cgst      = to_float(data.get("cgst"))
    sgst      = to_float(data.get("sgst"))
    igst      = to_float(data.get("igst"))
    total_tax = to_float(data.get("total_tax"))
    total     = to_float(data.get("total"))
    subtotal, cgst, sgst, igst, total_tax, total = fix_amounts(
        subtotal, cgst, sgst, igst, total_tax, total)
    hsn_inv = clean(data.get("hsn_sac"))
    seen = set(); lines = []
    for l in (data.get("lines") or []):
        if not isinstance(l, dict): continue
        item = clean(l.get("item") or l.get("description") or "")
        qty  = to_float(l.get("quantity") or l.get("qty"))
        rate = to_float(l.get("rate") or l.get("unit_price"))
        hsn  = clean(l.get("hsn_sac") or l.get("hsn") or "") or hsn_inv
        amt  = to_float(l.get("amount") or l.get("total"))
        if not item: continue
        key = (item.lower()[:40], round(amt))
        if key in seen: continue
        seen.add(key)
        if amt==0 and qty==0 and rate==0 and len(item)<6: continue
        lines.append({"item":item,"quantity":qty,"rate":rate,
                      "hsn_sac":hsn,"amount":amt})
    for line in lines:
        if not line["hsn_sac"] and hsn_inv:
            line["hsn_sac"] = hsn_inv
    hw = data.get("handwritten_references") or ""
    if isinstance(hw, list): hw = ", ".join(str(x) for x in hw if x)
    return {
        "filename":    filename,
        "vendor":      clean(data.get("vendor")),
        "invoice_number": clean(data.get("invoice_number")),
        "invoice_date":   clean(data.get("invoice_date")),
        "po_number":      clean(data.get("po_number")),
        "gstin":          clean(data.get("gstin")),
        "hsn_sac":        hsn_inv,
        "vendor_code":    clean(data.get("vendor_code")),
        "currency":       clean(data.get("currency")) or "INR",
        "subtotal":   round(subtotal,2),
        "cgst":       round(cgst,2),
        "sgst":       round(sgst,2),
        "igst":       round(igst,2),
        "total_tax":  round(total_tax,2),
        "total":      round(total,2),
        "grn_number": clean(data.get("grn_number") or ""),
        "handwritten_references": clean(hw),
        "confidence": round(max(0,min(1,to_float(data.get("confidence"))))*100,1),
    }, lines

def create_excel(header, lines, out_folder):
    Path(out_folder).mkdir(parents=True, exist_ok=True)
    fname = re.sub(r'[<>:"/\\|?*]',"_",
                   header["invoice_number"] or
                   Path(header["filename"]).stem).strip(" ._") + ".xlsx"
    path  = Path(out_folder)/fname
    wb    = Workbook()
    alt   = PatternFill("solid",fgColor="EBF3FB")

    def shdr(ws, row, color):
        fill=PatternFill("solid",fgColor=color)
        font=Font(bold=True,color="FFFFFF",size=10)
        for c in ws[row]:
            c.fill=fill; c.font=font
            c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

    ws1 = wb.active; ws1.title="Invoice Data"
    ws1.append(["Field","Value"]); shdr(ws1,1,"1F4E79")
    for i,(f,v) in enumerate([
        ("File Name",header["filename"]),
        ("Vendor",header["vendor"]),
        ("Invoice Number",header["invoice_number"]),
        ("Invoice Date",header["invoice_date"]),
        ("PO Number",header["po_number"]),
        ("Vendor GSTIN",header["gstin"]),
        ("HSN / SAC",header["hsn_sac"]),
        ("Vendor Code",header["vendor_code"]),
        ("Currency",header["currency"]),
        ("Subtotal",header["subtotal"]),
        ("CGST",header["cgst"]),
        ("SGST",header["sgst"]),
        ("IGST",header["igst"]),
        ("Total Tax",header["total_tax"]),
        ("Invoice Total",header["total"]),
        ("GRN / Inward Number",header["grn_number"]),
        ("Handwritten References",header["handwritten_references"]),
        ("AI Confidence %",header["confidence"]),
    ]):
        ws1.append([f,v])
        if i%2==1:
            for c in ws1[i+2]: c.fill=alt
    ws1.column_dimensions["A"].width=28
    ws1.column_dimensions["B"].width=65
    ws1.freeze_panes="A2"

    ws2 = wb.create_sheet("Line Items")
    ws2.append(["Sl.No","Item / Description","Quantity","Rate","HSN/SAC","Amount"])
    shdr(ws2,1,"145A32")
    for i,l in enumerate(lines,1):
        ws2.append([i,l["item"],l["quantity"],l["rate"],l["hsn_sac"],l["amount"]])
        if i%2==0:
            for c in ws2[i+1]: c.fill=alt
    for i,w in enumerate([8,60,12,15,18,18],1):
        ws2.column_dimensions[get_column_letter(i)].width=w
    ws2.freeze_panes="A2"
    wb.save(path)
    print(f"  📄 Excel: {fname}")


# ── Main ──────────────────────────────────────────────────────

def main():
    folder = Path(PDF_FOLDER)
    out    = Path(OUTPUT_FOLDER)

    if not folder.exists():
        print(f"ERROR: Folder not found: {folder}"); sys.exit(1)

    seen=set(); pdfs=[]
    for p in sorted(folder.iterdir()):
        if p.is_file() and p.suffix.lower()==".pdf":
            if p.name.lower() not in seen:
                seen.add(p.name.lower()); pdfs.append(p)

    if not pdfs: print("ERROR: No PDFs found"); sys.exit(1)

    out.mkdir(parents=True, exist_ok=True)

    print("="*60)
    print("KTFL Invoice Extractor — LangChain + Ollama + PostgreSQL")
    print(f"PDFs   : {len(pdfs)} | Model: {OLLAMA_MODEL}")
    print(f"DB     : {DB_URL.split('@')[-1]}")
    print("="*60)

    # DB connect
    print("\nConnecting to DB...")
    try:
        engine, Session = setup_db()
    except Exception as e:
        print(f"ERROR: DB failed: {e}"); sys.exit(1)

    # Ollama check
    try:
        r = httpx.get(f"{OLLAMA_URL}/api/tags",timeout=10,trust_env=False)
        models = [m["name"] for m in r.json().get("models",[])]
        if OLLAMA_MODEL not in models:
            print(f"ERROR: {OLLAMA_MODEL} not installed."); sys.exit(1)
        print(f"Ollama : OK\n")
    except Exception as e:
        print(f"ERROR: Ollama failed: {e}"); sys.exit(1)

    llm = create_llm()
    stats = {"pushed":0, "duplicate":0, "error":0}

    for idx, pdf in enumerate(pdfs, 1):
        t0 = time.time()
        print(f"\n[{idx}/{len(pdfs)}] {pdf.name}")

        try:
            # Extract
            images   = pdf_to_images(pdf)
            pdf_text = get_pdf_text(pdf)
            data     = invoke(llm, images, pdf_text)
            header, lines = normalize(data, pdf.name)

            print(f"  Invoice : {header['invoice_number'] or '?'}")
            print(f"  Vendor  : {header['vendor'] or '?'}")
            print(f"  Total   : {header['total']}  GRN: {header['grn_number'] or '-'}")

            with Session() as session:
                # Duplicate check via invoice_master
                dup, reason = is_duplicate(session, header["invoice_number"])
                print(f"  Check   : {reason}")

                if dup:
                    print(f"  ⚠️  SKIPPED — Duplicate")

                    # Log duplicate attempt in audit_logs
                    session.add(AuditLog(
                        invoice_id = None,
                        action     = "DUPLICATE_SKIPPED",
                        detail     = {
                            "file"           : pdf.name,
                            "invoice_number" : header["invoice_number"],
                            "vendor"         : header["vendor"],
                            "reason"         : reason,
                        }
                    ))
                    session.commit()
                    stats["duplicate"] += 1
                    create_excel(header, lines, out)
                    continue

                # Push to DB
                inv_id  = push_to_db(session, header, lines, pdf.name)
                elapsed = round(time.time()-t0, 1)
                print(f"  ✅ DB pushed — id={inv_id}  ({elapsed}s)")
                stats["pushed"] += 1

            # Save Excel
            create_excel(header, lines, out)

        except Exception as e:
            elapsed = round(time.time()-t0, 1)
            print(f"  ❌ ERROR: {e}  ({elapsed}s)")

            # Log error in audit_logs
            try:
                with Session() as session:
                    session.add(AuditLog(
                        invoice_id = None,
                        action     = "EXTRACTION_ERROR",
                        detail     = {
                            "file"  : pdf.name,
                            "error" : str(e),
                        }
                    ))
                    session.commit()
            except: pass

            stats["error"] += 1

        if idx < len(pdfs): time.sleep(0.5)

    print("\n" + "="*60)
    print("DONE")
    print(f"  ✅ Pushed to DB  : {stats['pushed']}")
    print(f"  ⚠️  Duplicates    : {stats['duplicate']} (skipped)")
    print(f"  ❌ Errors        : {stats['error']}")
    print(f"  📁 Excel folder  : {out}")
    print("="*60)


if __name__ == "__main__":
    main()
